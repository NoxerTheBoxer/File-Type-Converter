"""Spreadsheet and data conversion between Excel (.xlsx), CSV, and JSON."""
import csv
import datetime as dt
import io
import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from .paths import atomic_write, unique_path

TARGETS = {"xlsx": ["csv", "json"], "csv": ["xlsx", "json"], "json": ["csv", "xlsx"]}
MAX_EXACT_DIGITS = 15  # Excel stores numbers with 15 significant digits; longer IDs must stay text

Sheets = dict[str, list[list]]  # sheet name -> rows; the first row is the header


# ---------- reading ----------

def read(src: Path, fmt: str) -> Sheets:
    if fmt == "xlsx":
        return _read_xlsx(src)
    if fmt == "csv":
        return {_sheet_name(src.stem): _read_csv(src)}
    return _read_json(src)


def _read_xlsx(src: Path) -> Sheets:
    wb = load_workbook(src, read_only=True, data_only=True)  # data_only: formula results, not formulas
    try:
        sheets = {}
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
            sheets[ws.title] = rows
        return sheets
    finally:
        wb.close()


def _read_csv(src: Path) -> list[list]:
    raw = src.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    # newline="" keeps line breaks inside quoted cells intact
    return [[_parse_number(v) for v in row] for row in csv.reader(io.StringIO(text, newline=""), dialect)]


def _read_json(src: Path) -> Sheets:
    data = json.loads(src.read_text(encoding="utf-8-sig"))
    if isinstance(data, dict) and data and all(isinstance(v, list) for v in data.values()):
        sheets: Sheets = {}
        for key, records in data.items():
            sheets[_unique_sheet_name(str(key), sheets)] = _records_to_rows(records)
        return sheets
    if isinstance(data, dict):
        data = [data]
    if not isinstance(data, list):
        raise ValueError("JSON must be a list of records (or an object of lists) to convert to a table")
    return {_sheet_name(src.stem): _records_to_rows(data)}


def _records_to_rows(records: list) -> list[list]:
    if all(isinstance(r, list) for r in records):
        return [[_cell(v) for v in r] for r in records]
    header: list[str] = []
    for r in records:
        for key in (r if isinstance(r, dict) else {"value": r}):
            if key not in header:
                header.append(key)
    rows = [header]
    for r in records:
        r = r if isinstance(r, dict) else {"value": r}
        rows.append([_cell(r.get(k)) for k in header])
    return rows


def _cell(v):
    return json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v


def _parse_number(v: str):
    """'42' -> 42 and '3.5' -> 3.5. Values that only look numeric stay text: '007', '+1 242',
    '1E5' (a code, not 100000), and IDs longer than Excel can store exactly."""
    s = v.strip()
    if not s or s[0] == "+" or any(c in s for c in "_eE") or (len(s) > 1 and s[0] == "0" and s[1] != "."):
        return v
    if sum(c.isdigit() for c in s) > MAX_EXACT_DIGITS:
        return v
    try:
        return int(s)
    except ValueError:
        pass
    try:
        f = float(s)
        return f if f == f and abs(f) != float("inf") else v  # reject nan/inf spelled as text
    except ValueError:
        return v


def _sheet_name(name: str) -> str:
    cleaned = "".join("_" if c in '[]:*?/\\' else c for c in name).strip() or "Sheet1"
    return cleaned[:31]  # Excel's limit


def _unique_sheet_name(name: str, taken) -> str:
    base = _sheet_name(name)
    candidate, n = base, 2
    while candidate.lower() in {t.lower() for t in taken}:  # Excel sheet names ignore case
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    return candidate


# ---------- writing ----------

def write(sheets: Sheets, folder: Path, stem: str, target: str) -> list[Path]:
    if target == "xlsx":
        return [_write_xlsx(sheets, unique_path(folder, stem, ".xlsx"))]
    if target == "json":
        return [_write_json(sheets, unique_path(folder, stem, ".json"))]
    # CSV holds one table, so a multi-sheet workbook becomes one file per sheet.
    written: list[Path] = []
    try:
        for name, rows in sheets.items():
            file_stem = stem if len(sheets) == 1 else f"{stem} - {_safe_filename(name)}"
            written.append(_write_csv(rows, unique_path(folder, file_stem, ".csv")))
    except Exception:
        for p in written:  # all or nothing: don't leave some sheets converted and others not
            p.unlink(missing_ok=True)
        raise
    return written


def _write_csv(rows: list[list], dst: Path) -> Path:
    with atomic_write(dst) as tmp:
        with open(tmp, "w", encoding="utf-8-sig", newline="") as f:  # BOM so Excel reads accents correctly
            csv.writer(f).writerows([[_csv_safe(_plain(v)) for v in r] for r in rows])
    return dst


def _write_json(sheets: Sheets, dst: Path) -> Path:
    tables = {name: _rows_to_records(rows) for name, rows in sheets.items()}
    payload = next(iter(tables.values())) if len(tables) == 1 else tables
    with atomic_write(dst) as tmp:
        tmp.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return dst


def _write_xlsx(sheets: Sheets, dst: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(_unique_sheet_name(name, wb.sheetnames))
        for row in rows:
            ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row])
        for row in ws.iter_rows():
            for cell in row:
                # Text that starts with "=" must stay text: never let imported data become a live formula.
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
    with atomic_write(dst) as tmp:
        wb.save(tmp)
    return dst


def _rows_to_records(rows: list[list]) -> list[dict]:
    if not rows:
        return []
    width = max(len(r) for r in rows)  # rows longer than the header still keep every value
    header, seen = [], set()
    for i in range(1, width + 1):
        h = rows[0][i - 1] if i <= len(rows[0]) else None
        name = str(h) if h not in (None, "") else f"Column {i}"
        while name in seen:
            name += "_"
        seen.add(name)
        header.append(name)
    records = []
    for r in rows[1:]:
        r = list(r) + [None] * (width - len(r))
        records.append({k: _plain(v) for k, v in zip(header, r)})
    return records


def _plain(v):
    """Make a cell value JSON/CSV friendly."""
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    return str(v)  # timedelta (durations), Decimal, and anything else openpyxl may return


_FORMULA_LIKE = re.compile(r"^[=@\t\r]|^[+\-].*[A-Za-z(]")


def _csv_safe(v):
    """Stop text like '=HYPERLINK(...)' from running as a formula when the CSV is opened in Excel.
    Numbers and number-like text such as '-5' or '+1 242' are left alone."""
    if isinstance(v, str) and _FORMULA_LIKE.match(v):
        return "'" + v
    return v


def _safe_filename(name: str) -> str:
    cleaned = "".join("_" if c in '<>:"/\\|?*' or ord(c) < 32 else c for c in name).strip(" .")
    return cleaned or "Sheet"
