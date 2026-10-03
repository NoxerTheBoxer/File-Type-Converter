import json

import pytest
from openpyxl import load_workbook

from converter import data, registry


def test_csv_to_xlsx_keeps_text_like_zip_codes(make_data):
    [out] = registry.convert(make_data("csv"), "csv", "xlsx")
    ws = load_workbook(out).active
    assert [c.value for c in ws[2]] == ["mango", 3, "00123"]
    assert [c.value for c in ws[3]] == ["guava", 4.5, "+1 242"]


def test_csv_to_json(make_data):
    [out] = registry.convert(make_data("csv"), "csv", "json")
    assert json.loads(out.read_text(encoding="utf-8"))[0] == {"name": "mango", "qty": 3, "zip": "00123"}


def test_xlsx_with_two_sheets_to_csv_makes_two_files(make_data):
    outs = registry.convert(make_data("xlsx"), "xlsx", "csv")
    assert sorted(p.name for p in outs) == ["sample - Fruit.csv", "sample - Veg.csv"]
    assert outs[0].read_text(encoding="utf-8-sig").splitlines()[1] == "mango,3"


def test_xlsx_to_json_keeps_sheets(make_data):
    [out] = registry.convert(make_data("xlsx"), "xlsx", "json")
    assert json.loads(out.read_text(encoding="utf-8")) == {
        "Fruit": [{"name": "mango", "qty": 3}],
        "Veg": [{"name": "okra", "qty": 7}],
    }


def test_json_with_mixed_keys_to_csv(make_data):
    [out] = registry.convert(make_data("json"), "json", "csv")
    lines = out.read_text(encoding="utf-8-sig").splitlines()
    assert lines[0] == "name,qty,tags"
    assert lines[2] == 'guava,,"[""sweet""]"'


def test_json_to_xlsx(make_data):
    [out] = registry.convert(make_data("json"), "json", "xlsx")
    assert load_workbook(out).active["A2"].value == "mango"


def test_imported_text_never_becomes_a_formula(tmp_path):
    src = tmp_path / "evil.csv"
    src.write_text('name\n=HYPERLINK("http://evil","click")\n', encoding="utf-8")
    [out] = registry.convert(src, "csv", "xlsx")
    cell = load_workbook(out).active["A2"]
    assert cell.data_type == "s"


def test_json_that_is_not_a_table_gives_clear_error(tmp_path):
    src = tmp_path / "num.json"
    src.write_text("42")
    with pytest.raises(ValueError, match="list of records"):
        registry.convert(src, "json", "csv")


def test_semicolon_csv(tmp_path):
    src = tmp_path / "eu.csv"
    src.write_text("a;b\n1;2\n", encoding="utf-8")
    assert data.read(src, "csv") == {"eu": [["a", "b"], [1, 2]]}


@pytest.mark.parametrize("raw,expected", [("42", 42), ("-3.5", -3.5), ("007", "007"), ("1_000", "1_000"), ("nan", "nan"), ("", "")])
def test_number_parsing(raw, expected):
    assert data._parse_number(raw) == expected
