"""Regression tests for bugs found in code review of v0.1."""
import datetime as dt
import json

from openpyxl import Workbook, load_workbook

from converter import registry
from converter.formats import detect


def test_line_breaks_inside_quoted_csv_cells_survive(tmp_path):
    src = tmp_path / "notes.csv"
    src.write_text('name,note\nmango,"line1\nline2"\n', encoding="utf-8", newline="\n")
    [out] = registry.convert(src, "csv", "json")
    assert json.loads(out.read_text(encoding="utf-8"))[0]["note"] == "line1\nline2"


def test_csv_starting_with_bm_is_not_a_bitmap(tmp_path):
    src = tmp_path / "cars.csv"
    src.write_text("BMW,price\nX5,60000\n", encoding="utf-8")
    assert detect(src) == "csv"


def test_duration_cells_convert_to_json(tmp_path):
    src = tmp_path / "times.xlsx"
    wb = Workbook()
    wb.active.append(["task", "took"])
    wb.active.append(["install", dt.timedelta(hours=1, minutes=30)])
    wb.save(src)
    [out] = registry.convert(src, "xlsx", "json")
    assert json.loads(out.read_text(encoding="utf-8")) == [{"task": "install", "took": "1:30:00"}]


def test_sheet_names_with_windows_illegal_chars_become_valid_filenames(tmp_path):
    src = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "Q1 | Sales"
    wb.active.append(["a"])
    wb.create_sheet('Q2 "final"').append(["b"])
    wb.save(src)
    outs = registry.convert(src, "xlsx", "csv")
    assert sorted(p.name for p in outs) == ["book - Q1 _ Sales.csv", "book - Q2 _final_.csv"]


def test_rows_longer_than_header_keep_every_value(tmp_path):
    src = tmp_path / "ragged.csv"
    src.write_text("a,b\n1,2,3\n", encoding="utf-8")
    [out] = registry.convert(src, "csv", "json")
    assert json.loads(out.read_text(encoding="utf-8")) == [{"a": 1, "b": 2, "Column 3": 3}]


def test_long_ids_and_codes_stay_exact(tmp_path):
    src = tmp_path / "ids.csv"
    src.write_text("card,code,amount\n4111111111111111,1E5,12.5\n", encoding="utf-8")
    [out] = registry.convert(src, "csv", "xlsx")
    ws = load_workbook(out).active
    assert [c.value for c in ws[2]] == ["4111111111111111", "1E5", 12.5]


def test_json_tables_with_colliding_sheet_names_are_all_kept(tmp_path):
    src = tmp_path / "multi.json"
    src.write_text(json.dumps({"Sales/2024": [{"x": 1}], "Sales_2024": [{"x": 2}]}), encoding="utf-8")
    [out] = registry.convert(src, "json", "xlsx")
    assert load_workbook(out).sheetnames == ["Sales_2024", "Sales_2024 (2)"]


def test_csv_output_neutralizes_formulas_but_keeps_plain_values(tmp_path):
    src = tmp_path / "in.json"
    src.write_text(json.dumps([
        {"v": '=HYPERLINK("http://evil","x")'}, {"v": "+cmd|calc"}, {"v": "@SUM(A1)"},
        {"v": "+1 242"}, {"v": "- item"}, {"v": -5},
    ]), encoding="utf-8")
    [out] = registry.convert(src, "json", "csv")
    lines = out.read_text(encoding="utf-8-sig").splitlines()[1:]
    # Excel evaluates "- item" as a formula too (it shows #NAME?), so it gets protected as well.
    assert lines == ['"\'=HYPERLINK(""http://evil"",""x"")"', "'+cmd|calc", "'@SUM(A1)", "+1 242", "'- item", "-5"]
